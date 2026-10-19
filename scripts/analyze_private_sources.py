"""Read-only local document extraction. Never uploads content or runs document code.

Run: uv run --with pypdf --with python-docx scripts/analyze_private_sources.py
Private evidence is written only under .local/private-source-analysis/.
PDF text extraction is not OCR; image-only pages require manual local review.
Workbook index and body captures are optional: missing evidence is reported, not fabricated.
The explicit --verify-report mode requires an existing curated local comparison report.
Sensitive vocabulary belongs only in the private comparison JSON's sensitive_terms list.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / ".local" / "private-source-analysis"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalized(value: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", value)).casefold()


def extract_pdf(path: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted and not reader.decrypt(""):
        return {"status": "encrypted_not_extracted", "units": []}
    units = []
    for page_number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        units.append({"locator": f"page:{page_number}", "text": text,
                      "text_chars": len(text.strip()),
                      "requires_local_visual_review": not bool(text.strip())})
    empty = [unit["locator"] for unit in units if not unit["text"].strip()]
    return {"status": "partial_text" if empty else "text_extracted",
            "page_count": len(units), "no_text_pages": empty, "units": units,
            "limitations": ["Text only: image text, QR targets and layout are not verified.",
                            "Empty text does not establish whether a page is scanned or blank."]}


def extract_docx(path: Path) -> dict:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    units = []
    paragraph_number = table_number = 0
    for node in document.element.body.iterchildren():
        if node.tag == qn("w:p"):
            paragraph_number += 1
            paragraph = Paragraph(node, document)
            units.append({"locator": f"paragraph:{paragraph_number}", "text": paragraph.text})
        elif node.tag == qn("w:tbl"):
            table_number += 1
            table = Table(node, document)
            for row_number, row in enumerate(table.rows, 1):
                for column_number, cell in enumerate(row.cells, 1):
                    units.append({"locator": f"table:{table_number}/row:{row_number}/cell:{column_number}",
                                  "text": cell.text})
    for section_number, section in enumerate(document.sections, 1):
        for part_name in ("header", "footer"):
            for number, paragraph in enumerate(getattr(section, part_name).paragraphs, 1):
                if paragraph.text.strip():
                    units.append({"locator": f"section:{section_number}/{part_name}/paragraph:{number}",
                                  "text": paragraph.text})
    return {"status": "text_extracted", "paragraph_count": paragraph_number,
            "table_count": table_number, "units": units,
            "limitations": ["DOCX paragraph locations are not rendered page numbers.",
                            "Embedded images, text boxes, tracked changes and QR targets are not verified.",
                            "Document links, macros and embedded instructions are never executed."]}


def workbook_rows(acceptance: dict) -> list:
    rows = []
    sheets = acceptance.get("sheets", [])
    if not isinstance(sheets, list):
        return rows
    for sheet in sheets:
        if not isinstance(sheet, dict):
            continue
        row_data = sheet.get("rows")
        structured = row_data.get("structured") if isinstance(row_data, dict) else None
        values = structured.get("values") if isinstance(structured, dict) else None
        if isinstance(values, list):
            rows.extend((sheet.get("name"), index, row) for index, row in enumerate(values, 1)
                        if index > 1 and isinstance(row, list))
    return rows


def optional_evidence(path: Path) -> tuple[dict | None, str]:
    """Unavailable optional evidence must not block document extraction."""
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except FileNotFoundError:
        return None, "missing"
    except (OSError, UnicodeError, ValueError):
        return None, "unreadable_or_invalid_json"
    if not isinstance(value, dict):
        return None, "invalid_json_object"
    return value, "available"


def run(root: Path = ROOT, output: Path | None = None) -> dict:
    output = output if output is not None else root / ".local" / "private-source-analysis"
    output.mkdir(parents=True, exist_ok=True)
    raw = root / "raw_data"
    paths = sorted((p for p in raw.rglob("*") if p.is_file()), key=lambda p: str(p).casefold()) if raw.is_dir() else []
    inventory = []
    extracted = []
    for number, path in enumerate(paths, 1):
        entry = {"id": f"S{number:02}", "path": str(path.resolve()),
                 "relative_path": path.relative_to(root).as_posix(),
                 "type": path.suffix.lower(), "bytes": path.stat().st_size,
                 "sha256_before": sha256(path)}
        try:
            if path.suffix.lower() == ".pdf":
                result = extract_pdf(path)
            elif path.suffix.lower() == ".docx":
                result = extract_docx(path)
            else:
                result = {"status": "unsupported_not_parsed", "units": []}
        except Exception as error:  # noqa: BLE001 - Preserve per-file parser errors and continue the inventory.
            result = {"status": "extraction_error", "error": f"{type(error).__name__}: {error}", "units": []}
        entry["sha256_after"] = sha256(path)
        entry["unchanged"] = entry["sha256_before"] == entry["sha256_after"]
        entry["status"] = result["status"]
        entry["text_chars"] = sum(len(unit["text"].strip()) for unit in result["units"])
        inventory.append(entry)
        extracted.append({**entry, **result})
    acceptance_path = root / ".local" / "excel-mcp-acceptance.json"
    acceptance, index_status = optional_evidence(acceptance_path)
    evidence_gaps = []
    if acceptance is None:
        evidence_gaps.append({"kind": "workbook_index", "path": str(acceptance_path), "status": index_status})
    rows = workbook_rows(acceptance or {})
    if acceptance is not None and not rows:
        index_status = "no_source_rows"
        evidence_gaps.append({"kind": "workbook_index", "path": str(acceptance_path), "status": index_status})
    pairings = []
    samples = {}
    sample_paths = {path.stem.removeprefix("sample-body-"): path
                    for path in (root / ".local").glob("sample-body-*.json")}
    expected_cells = {f"{column}{row_index}" for _, row_index, row in rows
                      for column_index, column in enumerate(("B", "C"), 1)
                      if len(row) > column_index and row[column_index]}
    for cell in sorted(set(sample_paths) | expected_cells):
        path = sample_paths.get(cell, root / ".local" / f"sample-body-{cell}.json")
        body, status = optional_evidence(path)
        if body is not None and not isinstance(body.get("text"), str):
            body, status = None, "missing_text_field"
        if body is None:
            evidence_gaps.append({"kind": "body_capture", "cell": cell, "path": str(path), "status": status})
            continue
        samples[cell] = {**body, "path": str(path), "sha256": sha256(path)}
    if not samples and not expected_cells:
        evidence_gaps.append({"kind": "body_captures", "status": "missing"})
    for sheet, row_index, row in rows:
        source_name = str(row[0]) if row else ""
        candidates = [entry for entry in inventory if normalized(Path(entry["path"]).name) == normalized(source_name)]
        cells = [f"{column}{row_index}" for column in ("B", "C")]
        sample_links = []
        for column_index, cell in enumerate(cells, 1):
            if cell in samples:
                sample_links.append({"cell": cell, "sample_path": samples[cell]["path"],
                                     "url_matches_workbook": len(row) > column_index and row[column_index] == samples[cell].get("url"),
                                     "role": "preview" if column_index == 1 else "recap"})
        pairings.append({"sheet": sheet, "source_cell": f"A{row_index}", "source_name": source_name,
                         "source_ids": [candidate["id"] for candidate in candidates],
                         "match_status": "exact_normalized_filename" if len(candidates) == 1 else "missing_or_ambiguous",
                         "samples": sample_links})
    report = {"generated_at": datetime.now(UTC).isoformat(), "raw_directory": str(raw),
              "raw_directory_exists": raw.is_dir(), "file_count": len(inventory),
              "type_counts": dict(Counter(entry["type"] for entry in inventory)),
              "all_source_hashes_unchanged": all(entry["unchanged"] for entry in inventory),
              "inventory": inventory, "pairings": pairings,
              "acceptance_path": str(acceptance_path),
              "acceptance_sha256": sha256(acceptance_path) if acceptance_path.is_file() else None,
              "index_status": index_status, "evidence_gaps": evidence_gaps,
              "comparison_evidence_status": "complete" if not evidence_gaps else "incomplete",
              "sample_evidence": {cell: {key: data.get(key) for key in ("path", "sha256", "capture_note")} for cell, data in samples.items()},
              "safety": {"uploads": False, "ocr": False, "macros_executed": False,
                         "instructions_in_documents_executed": False, "source_writes": False}}
    baseline_path = output / "source-hash-baseline.json"
    current_hashes = {entry["relative_path"]: entry["sha256_before"] for entry in inventory}
    if not baseline_path.exists():
        baseline_path.write_text(json.dumps(current_hashes, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    baseline = json.loads(baseline_path.read_text(encoding="utf-8"))
    report["matches_first_recorded_source_hashes"] = baseline == current_hashes
    report["validation"] = {
        "raw_directory_exists": raw.is_dir(),
        "expected_source_rows_present": all(len(p["source_ids"]) == 1 for p in pairings) if pairings else None,
        "samples_match_workbook_urls": (all(s["url_matches_workbook"] for p in pairings for s in p["samples"])
                                        if any(p["samples"] for p in pairings) else None),
        "all_sources_extracted": bool(inventory) and all(e["status"] == "text_extracted" for e in inventory),
        "all_source_hashes_unchanged": report["all_source_hashes_unchanged"],
        "baseline_hashes_unchanged": report["matches_first_recorded_source_hashes"],
        "all_extracted_units_have_locators": all(u.get("locator") for e in extracted for u in e["units"]),
    }
    # Missing optional evidence is explicitly unknown, not a failed extraction.
    report["success"] = all(value is not False for value in report["validation"].values())
    for name, data in (("inventory.json", report), ("extractions.json", extracted)):
        (output / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    # Console output is deliberately counts/status only, never document content.
    print(json.dumps({"file_count": len(inventory), "type_counts": report["type_counts"],
                      "raw_directory_exists": raw.is_dir(), "all_source_hashes_unchanged": report["all_source_hashes_unchanged"],
                      "statuses": dict(Counter(entry["status"] for entry in inventory)),
                      "index_status": index_status, "evidence_gap_count": len(evidence_gaps),
                      "comparison_evidence_status": report["comparison_evidence_status"],
                      "success": report["success"]}, ensure_ascii=False))
    return report


def self_test() -> None:
    """Synthetic local fixtures only; no private material leaves this process."""
    from tempfile import TemporaryDirectory

    from docx import Document
    from pypdf import PdfWriter

    fixtures = OUTPUT / "self-test"
    fixtures.mkdir(parents=True, exist_ok=True)
    docx_path = fixtures / "synthetic.docx"
    document = Document()
    document.add_paragraph("LOCAL SYNTHETIC PARAGRAPH")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "LOCAL SYNTHETIC CELL"
    document.save(docx_path)
    before = sha256(docx_path)
    result = extract_docx(docx_path)
    assert any(u["locator"] == "paragraph:1" and u["text"] == "LOCAL SYNTHETIC PARAGRAPH" for u in result["units"])
    assert any(u["locator"] == "table:1/row:1/cell:1" and u["text"] == "LOCAL SYNTHETIC CELL" for u in result["units"])
    assert sha256(docx_path) == before
    pdf_path = fixtures / "synthetic-blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with pdf_path.open("wb") as target:
        writer.write(target)
    before = sha256(pdf_path)
    result = extract_pdf(pdf_path)
    assert result["no_text_pages"] == ["page:1"]
    assert result["units"][0]["requires_local_visual_review"]
    assert sha256(pdf_path) == before
    assert normalized(" A　B.docx ") == normalized("ab.docx")
    with TemporaryDirectory(prefix="standalone-", dir=fixtures) as directory:
        standalone = Path(directory)
        (standalone / "raw_data").mkdir()
        document.save(standalone / "raw_data" / "synthetic.docx")
        no_captures = run(root=standalone)
        assert no_captures["success"] and no_captures["file_count"] == 1
        assert no_captures["index_status"] == "missing"
        assert no_captures["validation"]["samples_match_workbook_urls"] is None
        assert {gap["kind"] for gap in no_captures["evidence_gaps"]} == {"workbook_index", "body_captures"}
        synthetic_index = {"sheets": [{"name": "Synthetic", "rows": {"structured": {"values": [
            ["source", "preview", "recap"], ["synthetic.docx", "synthetic-preview", "synthetic-recap"]]}}}]}
        local = standalone / ".local"
        (local / "excel-mcp-acceptance.json").write_text(json.dumps(synthetic_index), encoding="utf-8")
        index_only = run(root=standalone)
        assert index_only["success"] and index_only["validation"]["expected_source_rows_present"]
        assert len(index_only["evidence_gaps"]) == 2
        assert all(gap["status"] == "missing" for gap in index_only["evidence_gaps"])
        (local / "sample-body-B2.json").write_text(json.dumps({"text": "SYNTHETIC BODY", "url": "synthetic-preview"}), encoding="utf-8")
        partial_captures = run(root=standalone)
        assert partial_captures["success"] and partial_captures["validation"]["samples_match_workbook_urls"]
        assert partial_captures["sample_evidence"]["B2"]["capture_note"] is None
        assert len(partial_captures["evidence_gaps"]) == 1
        (local / "excel-mcp-acceptance.json").write_text("not json", encoding="utf-8")
        invalid_index = run(root=standalone)
        assert invalid_index["success"] and invalid_index["index_status"] == "unreadable_or_invalid_json"
        assert invalid_index["comparison_evidence_status"] == "incomplete"
    outcome = {"status": "passed", "checks": ["docx_paragraph_locator", "docx_table_cell_locator", "blank_pdf_no_ocr",
               "source_hash_unchanged", "unicode_filename_normalization", "standalone_without_index_or_bodies",
               "index_without_bodies", "partial_capture_without_capture_note", "invalid_optional_index"]}
    (OUTPUT / "self-test.json").write_text(json.dumps(outcome, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(outcome))


def verify_report() -> bool:
    """Verify curated evidence anchors and re-hash all originals without changing them."""
    inventory = json.loads((OUTPUT / "inventory.json").read_text(encoding="utf-8"))
    extractions = json.loads((OUTPUT / "extractions.json").read_text(encoding="utf-8"))
    comparison = json.loads((OUTPUT / "comparison.json").read_text(encoding="utf-8"))
    sources = {entry["id"]: entry for entry in extractions}
    failures = []
    finding_count = anchor_count = 0
    for row in comparison["rows"]:
        source = sources[row["source_id"]]
        locators = {unit["locator"] for unit in source["units"]}
        for finding in row["findings"]:
            finding_count += 1
            for span in finding["source_locator"].split(";"):
                match = re.fullmatch(r"\s*(page|paragraph):(\d+)(?:\.\.(\d+))?\s*", span)
                if not match:
                    failures.append(f"Invalid locator in {row['source_id']}: {span}")
                    continue
                kind, start, end = match.groups()
                for number in range(int(start), int(end or start) + 1):
                    if f"{kind}:{number}" not in locators:
                        failures.append(f"Missing locator: {row['source_id']} {kind}:{number}")
            if "sample" in finding:
                sample = json.loads((ROOT / ".local" / finding["sample"]).read_text(encoding="utf-8-sig"))
                anchor_count += 1
                if normalized(finding["sample_anchor"]) not in normalized(sample["text"]):
                    failures.append(f"Missing sample anchor: {row['source_id']} {finding['field']}")
    for entry in inventory["inventory"]:
        if sha256(Path(entry["path"])) != entry["sha256_before"]:
            failures.append(f"Changed source hash: {entry['id']}")
    public_text = (ROOT / "docs" / "research" / "source-comparison.md").read_text(encoding="utf-8")
    forbidden = [Path(entry["path"]).name for entry in inventory["inventory"]]
    # Sensitive vocabulary is private data, never a literal in this public script.
    forbidden.extend(value for value in comparison.get("sensitive_terms", []) if isinstance(value, str) and value)
    script_text = Path(__file__).read_text(encoding="utf-8")
    for value in forbidden:
        if value in public_text:
            failures.append("Private filename or curated personal name in public summary")
        if value in script_text:
            failures.append("Private filename or curated personal name in public script")
    if re.search(r"https?://|\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日|\d{4}-\d{2}-\d{2}", public_text):
        failures.append("URL or exact date in public summary")
    for data in inventory["sample_evidence"].values():
        if sha256(Path(data["path"])) != data["sha256"]:
            failures.append("Changed sample capture hash")
    outcome = {"status": "passed" if not failures else "failed", "source_count": len(sources),
               "finding_count": finding_count, "sample_anchor_count": anchor_count,
               "checks": ["curated_source_locators", "curated_sample_anchors", "source_hashes_rechecked",
                          "sample_hashes_rechecked", "public_summary_filename_name_url_date_screen",
                          "public_script_private_literal_screen"],
               "limitations": "Anchor existence verifies traceability, not the semantic correctness of every editorial finding.",
               "failures": failures}
    (OUTPUT / "evidence-verification.json").write_text(json.dumps(outcome, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: value for key, value in outcome.items() if key not in ("failures", "limitations")}))
    return not failures


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--verify-report", action="store_true")
    args = parser.parse_args()
    if args.self_test:
        self_test()
    elif args.verify_report:
        raise SystemExit(0 if verify_report() else 1)
    else:
        report = run()
        raise SystemExit(0 if report["success"] else 1)
